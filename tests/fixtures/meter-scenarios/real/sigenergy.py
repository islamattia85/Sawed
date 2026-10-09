"""Read a Sigenergy (mySigen) energy export, CSV or Excel, into half hours.

    read_sigenergy(path)  -> dict(rows={utc_start: {gen, load, imp, exp, bch, bdis}}, minutes, found, unit)
    to_half_hours(sig, esb)  -> {utc_half_hour_start: {...}}, with the timing checked against the ESB file

The columns are found by their headers, not their position, as the export's
layout is not fixed: generation (PV, solar, production, yield), home use (load,
consumption), bought (import, purchase, from grid), sold (export, feed-in, to
grid), battery charge and discharge. Values in kWh are energy per row; in kW or
W they are power, turned into energy by the row's length. Times are Irish local
time unless the header says UTC.

Whether a row is stamped at the start or the end of its interval is not
assumed: both are tried, and the one whose sales line up with the ESB file's
(which records every unit sold, half hour by half hour) is kept. The match is
reported, so a file that does not line up is caught before it is used.
"""
import csv, datetime as dt, io, os
from zoneinfo import ZoneInfo

TZ = ZoneInfo('Europe/Dublin')
KEYS = {  # first match wins; 'discharg' before 'charg' so a discharge column is not read as charge
    'bdis': ('discharg',),
    'bch': ('battery charg', 'charge energy', 'charging'),
    'gen': ('pv ', 'pv(', 'pv_', 'solar', 'generation', 'production', 'yield', 'pv'),
    'exp': ('export', 'feed-in', 'feed in', 'feedin', 'to grid', 'sell', 'sold'),
    'imp': ('import', 'purchas', 'from grid', 'buy', 'bought'),
    'load': ('load', 'consumption', 'home', 'house'),
}
SKIP = ('soc', 'state of charge', 'voltage', 'current', 'temperature', 'frequency', 'total', 'cumulative', 'lifetime')


def _rows(path):
    if path.lower().endswith(('.xlsx', '.xlsm')):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        return [[('' if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
    text = open(path, 'rb').read().decode('utf-8-sig', errors='replace')
    dialect = csv.Sniffer().sniff(text[:4000], delimiters=',;\t')
    return list(csv.reader(io.StringIO(text), dialect))


def _time(v):
    if isinstance(v, dt.datetime): return v
    s = str(v).strip()
    for f in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y/%m/%d %H:%M:%S', '%Y/%m/%d %H:%M', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M',
              '%d-%m-%Y %H:%M:%S', '%d-%m-%Y %H:%M', '%Y-%m-%dT%H:%M:%S', '%Y-%m-%dT%H:%M'):
        try: return dt.datetime.strptime(s, f)
        except ValueError: pass
    return None


def _num(v):
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace(' ', '')
    if not s or s in ('-', '--'): return 0.0
    if s.count(',') == 1 and '.' not in s: s = s.replace(',', '.')
    try: return float(s.replace(',', ''))
    except ValueError: return None


def read_sigenergy(path):
    rows = _rows(path)
    # The header: the first row naming a time and at least two of the quantities.
    hi = next((i for i, r in enumerate(rows[:30]) if any('time' in str(c).lower() or 'date' in str(c).lower() for c in r)
               and sum(any(k in str(c).lower() for ks in KEYS.values() for k in ks) for c in r) >= 2), None)
    if hi is None: raise ValueError('no header row with a time and generation/use columns')
    head = [str(c).strip() for c in rows[hi]]
    low = [h.lower() for h in head]
    tcol = next(i for i, h in enumerate(low) if 'time' in h or 'date' in h)
    found, taken = {}, set()
    for q, keys in KEYS.items():
        for i, h in enumerate(low):
            if i == tcol or i in taken or any(x in h for x in SKIP): continue
            if q == 'bch' and 'discharg' in h: continue
            if any(k in h for k in keys): found[q] = i; taken.add(i); break
    if 'gen' not in found: raise ValueError(f'no generation column among: {head}')
    utc = 'utc' in low[tcol]
    unit = {q: ('kwh' if 'kwh' in low[i] else 'w' if '(w)' in low[i] or low[i].endswith(' w') else 'kw' if 'kw' in low[i] else 'kwh') for q, i in found.items()}
    stamps, vals = [], []
    for r in rows[hi + 1:]:
        if tcol >= len(r): continue
        t = _time(r[tcol])
        if not t: continue
        v = {q: _num(r[i]) if i < len(r) else None for q, i in found.items()}
        if all(x is None for x in v.values()): continue
        stamps.append(t); vals.append(v)
    if len(stamps) < 48: raise ValueError(f'only {len(stamps)} rows with a time')
    order = sorted(range(len(stamps)), key=lambda k: stamps[k])
    stamps = [stamps[k] for k in order]; vals = [vals[k] for k in order]
    gaps = sorted((stamps[k + 1] - stamps[k]).total_seconds() / 60 for k in range(len(stamps) - 1))
    minutes = gaps[len(gaps) // 2]
    if minutes > 60: raise ValueError(f'rows are {minutes:.0f} minutes apart: too coarse to set the hours (daily totals cannot price night rates)')
    out, fold = {}, {}
    for t, v in zip(stamps, vals):
        if utc: u = t.replace(tzinfo=dt.timezone.utc)
        else:
            n = fold.get(t, 0); fold[t] = n + 1
            u = t.replace(tzinfo=TZ, fold=min(n, 1)).astimezone(dt.timezone.utc)
        e = {}
        for q, x in v.items():
            if x is None: x = 0.0
            e[q] = x if unit[q] == 'kwh' else x * minutes / 60 / (1000 if unit[q] == 'w' else 1)
        out[u] = e
    return dict(rows=out, minutes=minutes, found={q: head[i] for q, i in found.items()}, unit=unit)


def _bucket(sig, shift):
    """Sum the rows into half hours; `shift` moves each stamp back (an end-of-interval stamp)."""
    hh = {}
    step = dt.timedelta(minutes=sig['minutes'])
    for u, e in sig['rows'].items():
        s = u - step * shift
        if sig['minutes'] <= 30:
            k = s.replace(minute=0 if s.minute < 30 else 30, second=0, microsecond=0)
            d = hh.setdefault(k, {}); [d.__setitem__(q, d.get(q, 0.0) + x) for q, x in e.items()]
        else:   # an hour: half to each of its half hours
            for k in (s, s + dt.timedelta(minutes=30)):
                d = hh.setdefault(k, {}); [d.__setitem__(q, d.get(q, 0.0) + x / 2) for q, x in e.items()]
    return hh


def to_half_hours(sig, esb):
    """Half hours, with the stamping (start or end of interval) that lines up with the ESB file's sales."""
    meter = {u: (i, x) for u, _, i, x in esb}
    best = None
    for shift in (0, 1):
        hh = _bucket(sig, shift)
        q = 'exp' if 'exp' in sig['found'] else 'gen'
        both = [k for k in hh if k in meter]
        if not both: continue
        a = [hh[k].get(q, 0.0) for k in both]; b = [meter[k][1] for k in both]
        ma, mb = sum(a) / len(a), sum(b) / len(b)
        cov = sum((x - ma) * (y - mb) for x, y in zip(a, b))
        va = sum((x - ma) ** 2 for x in a) ** 0.5; vb = sum((y - mb) ** 2 for y in b) ** 0.5
        r = cov / (va * vb) if va and vb else 0.0
        if not best or r > best[0]: best = (r, shift, hh, len(both))
    if not best: raise ValueError('the Sigenergy file and the meter file share no half hours')
    r, shift, hh, n = best
    return hh, dict(stamped='end of interval' if shift else 'start of interval', match=round(r, 3), half_hours_shared=n)
